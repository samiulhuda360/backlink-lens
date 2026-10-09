"""Read a backlink export (CSV, TSV or XLSX) into a header row plus a list of row dicts."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 50_000


@dataclass
class Table:
    name: str
    headers: list[str]
    rows: list[dict[str, str]]


class ReadError(ValueError):
    """The file could not be read as a table."""


def read_table(name: str, data: bytes) -> Table:
    """Parse the bytes of a CSV/TSV/XLSX file. The first non-empty row is the header."""
    suffix = Path(name).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(name, data)
    if suffix in {".csv", ".tsv", ".txt", ""}:
        return _read_csv(name, data)
    raise ReadError(f"{name}: unsupported file type '{suffix}' (use .csv, .tsv or .xlsx)")


def read_file(path: Path) -> Table:
    return read_table(path.name, path.read_bytes())


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" not in text:
            return text
    raise ReadError("file is not text in a known encoding")


def _read_csv(name: str, data: bytes) -> Table:
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _build(name, list(reader))


def _read_xlsx(name: str, data: bytes) -> Table:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ReadError(f"{name}: not a readable workbook ({exc.__class__.__name__})") from exc
    sheet = book.worksheets[0]
    records: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        records.append(["" if value is None else str(value) for value in row])
        if len(records) > MAX_ROWS:
            break
    book.close()
    return _build(name, records)


def _build(name: str, records: list[list[str]]) -> Table:
    """The first non-empty row is the header; every row after it counts, blank ones included."""
    start = next((i for i, row in enumerate(records) if any(cell.strip() for cell in row)), None)
    if start is None:
        raise ReadError(f"{name}: the file is empty")
    headers = [cell.strip() for cell in records[start]]
    while headers and not headers[-1]:
        headers.pop()
    rows: list[dict[str, str]] = []
    for record in records[start + 1 : start + 1 + MAX_ROWS]:
        padded = list(record) + [""] * (len(headers) - len(record))
        rows.append({header: padded[i].strip() for i, header in enumerate(headers) if header})
    return Table(name=name, headers=[h for h in headers if h], rows=rows)
