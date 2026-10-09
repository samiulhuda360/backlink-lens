"""Write the comparison workbook: summary bands, one clean sheet per competitor, flags and the change log."""

from __future__ import annotations

import csv
import io
from dataclasses import fields as dataclass_fields

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .bands import Comparison
from .cleaning import CleanResult, Row
from .columns import FIELDS, Mapping

CellValue = str | int | float | None

HEADER_FILL = PatternFill("solid", fgColor="1F3A3D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
AVERAGE_FONT = Font(italic=True)
ROW_FIELDS = [f.name for f in dataclass_fields(Row)]


def build_workbook(results: list[CleanResult], comparison: Comparison, mappings: dict[str, Mapping], notes: str = "") -> bytes:
    book = Workbook()
    book.remove(book.worksheets[0])
    _write_summary(book.create_sheet("Summary"), comparison, notes)
    used: set[str] = set()
    for result in results:
        title = _sheet_name(result.competitor)
        while title in used:
            title = f"{title[:28]}({len(used) + 1})"
        used.add(title)
        _write_rows(book.create_sheet(title), result)
    _write_flags(book.create_sheet("Flags"), results)
    _write_changes(book.create_sheet("Change log"), results)
    _write_mapping(book.create_sheet("Column mapping"), mappings)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def changes_csv(results: list[CleanResult]) -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(["file", "competitor", "row", "field", "before", "after", "reason"])
    for result in results:
        for change in result.changes:
            writer.writerow([result.name, result.competitor, change.row, change.field, change.before, change.after, change.reason])
    return out.getvalue()


def _header(sheet: Worksheet, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width


def _write_summary(sheet: Worksheet, comparison: Comparison, notes: str) -> None:
    row = 1
    for title, labels, attr, average in (
        ("Rating", comparison.dr_labels, "dr_counts", comparison.dr_average),
        ("Referring domains", comparison.rd_labels, "rd_counts", comparison.rd_average),
    ):
        _header(sheet, row, ["Competitor", *[f"{title} {label}" for label in labels], "Total"])
        for profile in comparison.profiles:
            row += 1
            counts = profile.dr_counts if attr == "dr_counts" else profile.rd_counts
            sheet.cell(row=row, column=1, value=profile.competitor)
            for col, count in enumerate(counts, start=2):
                sheet.cell(row=row, column=col, value=count)
            sheet.cell(row=row, column=len(counts) + 2, value=sum(counts))
        row += 1
        sheet.cell(row=row, column=1, value="Average").font = AVERAGE_FONT
        for col, value in enumerate(average, start=2):
            sheet.cell(row=row, column=col, value=value).font = AVERAGE_FONT
        row += 3
    _header(sheet, row, ["Competitor", "Clean links", "Median rating", "Links rated 50+", "Rows without metrics"])
    for profile in comparison.profiles:
        row += 1
        stats: list[CellValue] = [profile.competitor, profile.rows, profile.median_rating, profile.strong_links, profile.unbanded]
        for col, stat in enumerate(stats, start=1):
            sheet.cell(row=row, column=col, value=stat)
    if notes:
        row += 2
        sheet.cell(row=row, column=1, value="Notes").font = Font(bold=True)
        for line in notes.splitlines():
            row += 1
            sheet.cell(row=row, column=1, value=line)
    _widths(sheet, [28] + [14] * 12)


def _write_rows(sheet: Worksheet, result: CleanResult) -> None:
    labels = {name: label for name, (label, _, _) in FIELDS.items()}
    _header(sheet, 1, ["Source row", *[labels.get(f, f) for f in ROW_FIELDS[1:]], "Flags"])
    flags_by_row: dict[int, list[str]] = {}
    for flag in result.flags:
        flags_by_row.setdefault(flag.row, []).append(flag.code)
    for r, row in enumerate(result.rows, start=2):
        values = [getattr(row, f) for f in ROW_FIELDS]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=r, column=col, value=value)
        sheet.cell(row=r, column=len(values) + 1, value=", ".join(flags_by_row.get(row.index, [])))
    _widths(sheet, [10, 30, 50, 14, 18, 50, 30, 12, 12, 30])
    sheet.freeze_panes = "A2"


def _write_flags(sheet: Worksheet, results: list[CleanResult]) -> None:
    _header(sheet, 1, ["File", "Competitor", "Source row", "Domain", "Check", "Severity", "What it means"])
    r = 1
    for result in results:
        for flag in result.flags:
            r += 1
            values: list[CellValue] = [result.name, result.competitor, flag.row, flag.domain, flag.code, flag.severity, flag.message]
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value)
    _widths(sheet, [26, 24, 10, 32, 16, 10, 80])
    sheet.freeze_panes = "A2"


def _write_changes(sheet: Worksheet, results: list[CleanResult]) -> None:
    _header(sheet, 1, ["File", "Competitor", "Source row", "Field", "Before", "After", "Reason"])
    r = 1
    for result in results:
        for change in result.changes:
            r += 1
            values: list[CellValue] = [result.name, result.competitor, change.row, change.field, change.before, change.after, change.reason]
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value)
    _widths(sheet, [26, 24, 10, 18, 40, 40, 40])
    sheet.freeze_panes = "A2"


def _write_mapping(sheet: Worksheet, mappings: dict[str, Mapping]) -> None:
    _header(sheet, 1, ["File", "Field", "Column in file", "Confidence", "Decided by"])
    r = 1
    for name, mapping in mappings.items():
        for field_name, header in mapping.fields.items():
            r += 1
            values: list[CellValue] = [name, FIELDS[field_name][0], header, mapping.confidence.get(field_name, 1.0), mapping.source.get(field_name, "")]
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value)
        for header in mapping.unmapped:
            r += 1
            sheet.cell(row=r, column=1, value=name)
            sheet.cell(row=r, column=2, value="(ignored)")
            sheet.cell(row=r, column=3, value=header)
    _widths(sheet, [26, 20, 30, 12, 12])


def _sheet_name(competitor: str) -> str:
    cleaned = "".join(ch if ch not in "[]:*?/\\" else "_" for ch in competitor)
    return cleaned[:31] or "competitor"
