import io

from openpyxl import load_workbook

from backlink_lens.bands import DR_BANDS, RD_BANDS, band_label, band_of, compare, profile
from backlink_lens.chart import band_chart
from backlink_lens.cleaning import Change, CleanResult, Row
from backlink_lens.columns import Mapping
from backlink_lens.report import build_workbook, changes_csv


def result(competitor: str, ratings: list[tuple[float, int]]) -> CleanResult:
    rows = [
        Row(index=i, referring_domain=f"d{i}.com", domain_rating=dr, referring_domains=rd, target_url=f"https://{competitor}/")
        for i, (dr, rd) in enumerate(ratings, start=1)
    ]
    return CleanResult(name=f"{competitor}.csv", competitor=competitor, rows=rows, source_rows=len(rows))


def test_band_boundaries() -> None:
    assert band_label(DR_BANDS[0]) == "0-10"
    assert band_label(DR_BANDS[-1]) == "91-100"
    assert band_label(RD_BANDS[-1]) == "1001+"
    assert band_of(10, DR_BANDS) == (0, 10)
    assert band_of(11, DR_BANDS) == (11, 20)
    assert band_of(0, RD_BANDS) is None  # zero referring domains is outside the 1-100 band
    assert band_of(None, DR_BANDS) is None


def test_profile_counts_and_median() -> None:
    p = profile(result("acme.com", [(5, 10), (55, 500), (95, 5000), (None, 10)]))  # type: ignore[list-item]
    assert p.dr_counts[0] == 1 and p.dr_counts[5] == 1 and p.dr_counts[9] == 1
    assert p.rd_counts[0] == 1 and p.rd_counts[4] == 1 and p.rd_counts[-1] == 1
    assert p.unbanded == 1
    assert p.median_rating == 55
    assert p.strong_links == 2


def test_compare_averages_and_gaps() -> None:
    strong = result("strong.com", [(75, 2000)] * 8)
    weak = result("weak.com", [(75, 2000)] * 1 + [(15, 50)] * 3)
    comparison = compare([strong, weak])
    assert comparison.dr_average[7] == 4.5
    gaps = comparison.gaps()
    assert gaps and gaps[0]["competitor"] == "weak.com" and gaps[0]["band"] == "rating 71-80"


def test_workbook_has_summary_clean_flags_and_changes() -> None:
    results = [result("acme.com", [(45, 1200), (70, 3000)]), result("other.net", [(20, 100)])]
    mapping = Mapping(fields={"domain_rating": "DR"}, confidence={"domain_rating": 1.0}, source={"domain_rating": "alias"}, unmapped=["Notes"])
    data = build_workbook(results, compare(results), {"acme.com.csv": mapping}, notes="line one\nline two")
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Summary", "acme.com", "other.net", "Flags", "Change log", "Column mapping"]
    summary = book["Summary"]
    assert summary["A1"].value == "Competitor"
    assert summary["A2"].value == "acme.com"
    assert summary["A4"].value == "Average"
    assert book["acme.com"]["B2"].value == "d1.com"
    assert book["Column mapping"]["C3"].value == "Notes"


def test_sheet_names_are_unique_and_valid() -> None:
    results = [result("same.com", [(1, 1)]), result("same.com", [(2, 2)]), result("a/b:c*very-long-competitor-name.example", [(3, 3)])]
    data = build_workbook(results, compare(results), {})
    names = load_workbook(io.BytesIO(data)).sheetnames
    assert len(set(names)) == len(names)
    assert all(len(n) <= 31 for n in names)


def test_changes_csv_has_header_and_rows() -> None:
    r = result("acme.com", [(1, 1)])
    r.changes.append(Change(1, "domain_rating", "1.0", "1", "number parsed"))
    text = changes_csv([r])
    assert text.splitlines()[0] == "file,competitor,row,field,before,after,reason"
    assert "number parsed" in text


def test_chart_geometry_fits_the_canvas() -> None:
    chart = band_chart(["a", "b", "c"], [("x", [1, 5, 0]), ("y", [3, 2, 7])], width=600, height=200)
    assert len(chart.bars) == 6
    assert all(0 <= b.x <= 600 and 0 <= b.y <= 200 and b.height >= 0 for b in chart.bars)
    assert chart.legend == [("x", chart.bars[0].colour), ("y", chart.bars[3].colour)]
    assert chart.gridlines[0][1] == "0"
